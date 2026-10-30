"""Genera el informe integrado de cierre; requiere python-docx y resultados congelados."""
import csv
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'docs/informes';OUT.mkdir(exist_ok=True)
DEST=OUT/'informe_comparacion_L0_V4_PARAAD_20261029.docx'
R=ROOT/'results/barcelona_paraad/examen_colab_2018_2019/resumen'

def rows(p):
    with p.open() as f:return list(csv.DictReader(f))

M=rows(R/'metricas.csv');CI=rows(R/'intervalos.csv');D=rows(R/'diferencias_pareadas.csv');FA=rows(R/'tasa_limpia.csv')
CO=rows(ROOT/'data/reference/barcelona_paraad/cohorte13.csv')
def select(data,**kw):return [r for r in data if all(str(r[k])==str(v) for k,v in kw.items())]
def met(y,delta,det,point='p99',estado='con_ataque'):return select(M,ano=y,delta=delta,detector=det,punto=point,estado=estado)[0]
def f(x):return f'{float(x):.4f}'
def pct(x):return f'{100*float(x):.4f}%'

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
sec.footer_distance=Inches(.25)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(6)
for n,size in [('Title',23),('Subtitle',12),('Heading 1',17),('Heading 2',12)]:
    s=doc.styles[n];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s.font.underline=False;s.paragraph_format.space_after=Pt(9);s.paragraph_format.keep_with_next=True
    if n=='Heading 1':s.paragraph_format.space_before=Pt(0)
    if n=='Heading 2':s.paragraph_format.space_before=Pt(9)
for n in ['List Bullet','Caption']:
    doc.styles[n].font.name='Arial';doc.styles[n].font.size=Pt(10 if n=='List Bullet' else 9)
    doc.styles[n].font.color.rgb=RGBColor(0,0,0)
    doc.styles[n].paragraph_format.space_after=Pt(5)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
rr=foot.add_run('Comparación de detectores · ');rr.font.name='Arial';rr.font.size=Pt(8)
for instr in ['PAGE','NUMPAGES']:
    if instr=='NUMPAGES':foot.add_run(' de ').font.size=Pt(8)
    r=foot.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),instr);r._r.addnext(field)
doc.core_properties.title='Comparación de L0 V4 y PARAAD en RAMA y Barcelona'
doc.core_properties.subject='Resultados logrados artefactos disponibles y alcance para la tesis'
doc.core_properties.author='Proyecto de tesis'
doc.core_properties.keywords='RAMA Barcelona PARAAD L0 V4 anomalías evaluación reproducibilidad'

def p(t,bold=False,style=None):
    q=doc.add_paragraph(style=style);r=q.add_run(t);r.bold=bold
    q.paragraph_format.widow_control=True
    return q
def h(t,level=1):return doc.add_paragraph(t,style=f'Heading {level}')
def page(title):doc.add_page_break();h(title)
def bullet(t):return p(t,style='List Bullet')
def caption(t):return p(t,style='Caption')

def table(head,data,widths=None,size=9.4):
    total=6.86
    if widths is None:widths=[total/len(head)]*len(head)
    t=doc.add_table(rows=1,cols=len(head));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for i,x in enumerate(head):t.rows[0].cells[i].text=str(x)
    for row in data:
        cells=t.add_row().cells
        for j,x in enumerate(row):cells[j].text=str(x)
    for i,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr();el=OxmlElement('w:cantSplit');trPr.append(el)
        if i==0:el=OxmlElement('w:tblHeader');trPr.append(el)
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr()
            marg=OxmlElement('w:tcMar')
            for side,value in [('top','80'),('bottom','80'),('left','90'),('right','90')]:
                z=OxmlElement('w:'+side);z.set(qn('w:w'),value);z.set(qn('w:type'),'dxa');marg.append(z)
            tcPr.append(marg)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','bottom','left','right']:
                z=OxmlElement('w:'+side);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'4');z.set(qn('w:color'),'D9D9D9');borders.append(z)
            tcPr.append(borders)
            fill='23384D' if i==0 else ('F2F5F8' if i%2==0 else 'FFFFFF')
            z=OxmlElement('w:shd');z.set(qn('w:fill'),fill);tcPr.append(z)
            for q in c.paragraphs:
                q.paragraph_format.space_after=Pt(0);q.paragraph_format.space_before=Pt(0);q.paragraph_format.line_spacing=1.05
                q.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in q.runs:
                    r.font.name='Arial';r.font.size=Pt(size);r.bold=(i==0)
                    r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
    p('').paragraph_format.space_after=Pt(0)
    return t

def f1rows(point):
    data=[]
    for y in [2018,2019]:
        for delta in [100,50,20,10,5]:data.append([y,f'{delta}%',*[f(met(y,delta,d,point)['f1']) for d in ['l0','v4','paraad']]])
    return data

# 1
p('Comparación de L0 V4 y PARAAD en RAMA y Barcelona',style='Title')
p('Resultados logrados artefactos disponibles y alcance para la tesis',style='Subtitle')
p('Informe para revisión con el tutor · Corte al 29 de octubre de 2026')
p('Se completaron las dos direcciones de la comparación: adaptar PARAAD a RAMA y entrenar L0, V4 y PARAAD sobre una cohorte propia de Barcelona. La fase proporcional de Barcelona ya tiene datos auditados, modelos reproducibles, calibración limpia y evaluación congelada en dos años.',bold=True)
p('En Barcelona, L0 obtuvo una ventaja respaldada por los intervalos de diferencia en ataques del 20% y 50% con p99, aunque su entrenamiento supervisado impide atribuirla sólo a la arquitectura. Entre los métodos sin etiquetas de ataque, PARAAD presentó mayores F1 puntuales que V4; la diferencia quedó más clara en 2019. Las anomalías pequeñas y el aumento temporal de falsas alarmas siguen siendo limitaciones.')
h('Qué quedó completado',2)
table(['Componente','Estado comprobado'],[
 ['Datos Barcelona','13 estaciones · 2014 a 2019 · procedencia y hashes'],
 ['Preparación común','Rejilla UTC · hora local · ventanas y máscaras'],
 ['Ataques proporcionales','70 calendarios · 150 variantes aplicadas'],
 ['Entrenamiento Barcelona','15 modelos definitivos · 64 ejecuciones científicas'],
 ['Calibración Barcelona','2017 limpio · 4,536 estación días comunes'],
 ['Examen Barcelona','2018 y 2019 · 100 variantes atacadas y dos limpias'],
 ['Infraestructura','Respaldos verificados · VMs liberadas · seguimientos cerrados']
],widths=[2.05,4.81])
p('Las tasas de falsas alarmas y las métricas se distinguen del entrenamiento. Los valores publicados en este informe provienen de modelos cargados de disco, sin reentrenamiento durante la evaluación. Los resultados de RAMA se recuperan de reportes ya cerrados, sin reabrir sus datos de examen para decidir parámetros.')
p('La parte principal del documento explica Barcelona. El antecedente RAMA se presenta como el otro brazo del estudio. El sesgo aditivo de PARAAD permanece como complemento registrado y pendiente de evaluación.')

# 2
page('Pregunta de investigación y comparación justa')
p('El objetivo es comparar detectores bajo condiciones compartidas y determinar si las conclusiones observadas en RAMA aparecen también al reconstruir el experimento con datos de Barcelona. Se conserva la idea de comparar “peras con peras”: mismos datos recibidos, calendarios, casos evaluables, unidades de decisión y reglas de calibración.')
table(['Dirección','Qué se hizo','Qué permite estudiar'],[
 ['PARAAD hacia RAMA','Modelo adaptado a la red y evaluación con detectores congelados','Comportamiento de un método publicado en nuestro escenario'],
 ['L0 y V4 hacia Barcelona','Modelos nuevos, junto con PARAAD, sobre 13 estaciones propias','Aplicación de la metodología en otro conjunto de datos']
],widths=[1.45,2.7,2.71])
h('Información que recibe cada detector',2)
table(['Método','Ajuste de pesos','Comparación que admite'],[
 ['L0 compartido','Observaciones alteradas y etiquetas de ataque con δ50','Referencia supervisada de detección'],
 ['V4','Concentraciones observadas como respuesta de regresión; sin etiquetas de ataque','Comparación principal con PARAAD sin supervisión de ataque'],
 ['PARAAD SA','Reconstrucción de observaciones sin inyección experimental','Comparación principal con V4']
],widths=[1.3,2.95,2.61])
p('Una respuesta de concentración para entrenar una regresión no equivale a una etiqueta de ataque. V4 aprende a predecir valores y PARAAD a reconstruirlos; ninguno recibe la indicación “esta muestra fue atacada” al ajustar pesos. L0 sí la recibe, separada de sus características de entrada.')
p('Los métodos conservan su representación: PARAAD reconstruye el bloque actual de la red, V4 excluye al objetivo de las entradas de predicción y L0 combina al objetivo con cuatro vecinas. Igualar la evaluación no convierte sus arquitecturas ni su información en idénticas.')
h('Alcance de la transferencia',2)
p('En Barcelona se entrenaron los modelos desde cero con datos locales. No se trasladaron los pesos ni los umbrales de RAMA. Por ello, el resultado respalda la aplicación del método y de la metodología en otro conjunto de datos, no una transferencia directa de pesos entre ciudades.')
p('El artículo original empleó ocho estaciones. Esta comparación usa una cohorte propia de 13 estaciones de la misma fuente pública; no se presenta como reproducción literal de sus tablas, partición o umbral móvil. [B1, L1]')

# 3
page('Datos de Barcelona y selección de estaciones')
p('Se descargaron series horarias de O3 de 2014 a 2019 del dataset público tasf-thgu de la Generalitat. La delimitación es la AMB administrativa. Entre las 14 candidatas, se conservaron las 13 aprobadas con disponibilidad de ajuste en 2014 y 2015: al menos 2,000 horas y seis meses con observaciones. Todas tienen datos en los 24 meses de ajuste. [B1, D1, D2]')
table(['Código EOI','Estación','Cobertura de ajuste'],[[x['codigo'],x['nombre'],f"{float(x['cobertura_ajuste_pct']):.2f}%"] for x in CO],widths=[1.0,4.56,1.3],size=9.2)
p('Observatori Fabra quedó fuera por no disponer de datos de ajuste en 2014 y 2015. La selección no usó puntajes de detección de 2018 o 2019. Las coordenadas se tomaron de las primeras observaciones de ajuste y el orden de estaciones quedó fijado por código EOI.')
h('Calidad y conservación de datos',2)
p('La cohorte aporta 660,451 valores horarios presentes. La auditoría no encontró duplicados de estación y fecha ni concentraciones negativas en la descarga. Se mantienen las unidades originales, µg/m³, sin convertir a ppb ni importar umbrales de contingencia de Ciudad de México.')
p('Después de apartar los días locales de cambio de reloj, la rejilla conserva 656,916 observaciones utilizables; otras 3,535 permanecen respaldadas y excluidas. “Limpio” significa sin inyección experimental: no certifica que las series originales carezcan de errores reales de sensor.')
p('Disponibilidad elevada no significa días siempre completos. Los faltantes permanecen como tales, con máscaras; no se fabrican lecturas para completar el día. La elegibilidad final depende también del contexto requerido por cada detector.')

# 4
page('Tiempo ventanas y separación temporal')
p('Se adoptó una convención operativa explícita: h01 representa el intervalo local de 00 a 01 y h24 el de 23 a 24, con timestamp al inicio. Las horas se localizan en Europe/Madrid y se convierten a UTC para medir continuidad; seno y coseno se calculan con la hora local para conservar el ciclo diario. La fecha completa no se añade como característica del modelo. [B2]')
p('Los días locales completos de transición estacional se excluyen para todas las estaciones, porque la fuente no permite desambiguar de forma certificada las horas repetidas. La rejilla conserva sus 23 o 25 intervalos físicos como excluidos. Una ventana que toca esas horas tampoco es válida; los huecos ordinarios se mantienen con máscara.')
table(['Método','Entrada por ejemplo','Contexto'],[
 ['PARAAD SA','13 sensores × 24 posiciones horarias','Bloque UTC completo de 00 a 23'],
 ['V4','24 horas × 26 características','Otras 12 estaciones, máscaras y hora; t−24 a t−1'],
 ['L0','24 horas × 17 características','Objetivo y cuatro vecinas; t−23 a t']
],widths=[1.2,2.0,3.66])
p('Los perfiles usan los 14 días locales previos a la misma hora, con respaldo histórico de esa hora y después global estrictamente pasado. El contexto anterior puede cruzar una etapa de forma causal; no se concatenan días separados como si fueran contiguos.')
table(['Etapa UTC','Uso','Estación días comunes'],[
 ['2014 y 2015','Ajuste inicial de pesos y escalas','9,149'],
 ['2016','Selección de épocas y configuración','4,538'],
 ['2014 a 2016','Ajuste definitivo desde cero','13,687'],
 ['2017','Calibración limpia de umbrales','4,536'],
 ['2018','Prueba principal','4,549'],
 ['2019','Prueba complementaria','4,567']
],widths=[1.25,3.91,1.7])
p('Cada estación día común requiere un bloque completo seguro, al menos 12 observaciones propias y al menos 12 horas evaluables para cada método. Se conserva el soporte horario propio de cada detector dentro del mismo universo diario. Los conteos de examen sólo describen disponibilidad; no seleccionan modelos.')

# 5
page('Ataques proporcionales y significado de las etiquetas')
p('El escenario principal traslada la alteración proporcional de RAMA. Cada estación libre puede iniciar un episodio diario con probabilidad p, sorteada por calendario entre 0.001 y 0.003. La duración es uniforme de uno a siete días y el signo se elige con igual probabilidad, fijo durante el episodio. Las 13 estaciones pueden ser atacadas. [B3]')
p('Con signo positivo se aplica x × (1 + δ); con signo negativo, x / (1 + δ). Se evalúan δ de 100%, 50%, 20%, 10% y 5%. Por ejemplo, con δ50 una lectura de 60 pasa a 90 o a 40. El caso negativo no usa x × (1 − δ).')
p('Sólo se alteran lecturas observadas. El valor se redondea al entero más cercano mediante floor(valor + 0.5), en µg/m³: las lecturas de ajuste eran enteras y se evita introducir decimales exclusivos del ataque. Los faltantes y los días excluidos por cambio de reloj permanecen intactos.')
table(['Uso','Calendarios','Deltas aplicados','Variantes'],[
 ['Ajuste inicial L0','20','50%','20'],['Validación L0','10','50%','10'],
 ['Ajuste definitivo L0','20','50%','20'],['Prueba 2018','10','Los cinco','50'],
 ['Complementaria 2019','10','Los cinco','50'],['Total','70','Según el uso','150']
],widths=[2.55,1.1,1.8,1.41])
p('Los cinco deltas comparten estaciones, fechas, duración y signo del mismo calendario. L0, V4 y PARAAD reciben la misma matriz alterada. Se permiten episodios simultáneos entre estaciones y episodios contiguos, pero no superposición en una estación. Los planes se sortean sin mirar concentraciones ni disponibilidad; no se vuelven a sortear episodios sin soporte.')
h('Dos etiquetas con propósitos distintos',2)
p('La etiqueta horaria de entrenamiento L0 indica una lectura efectivamente cambiada tras redondear. La etiqueta diaria del examen indica pertenencia al episodio programado. Por eso una hora puede estar programada y conservar su valor, especialmente con concentraciones bajas y δ pequeño. La máscara diaria no se redefine según la intensidad.')
p('Las inyecciones se aplican antes de construir entradas. Se recalculan perfiles, relativas y columnas de vecinas sobre la serie alterada, conservando las escalas limpias. Así también se conserva la contaminación del contexto después del episodio; no se facilita la detección usando un pasado artificialmente limpio dentro de la variante.')

# 6
page('Entrenamiento y explicación de las épocas elegidas')
p('Los 15 modelos definitivos de Barcelona son nuevos: un L0 compartido, 13 V4 por objetivo y un PARAAD SA. La selección se realizó sólo con 2016, y el definitivo se entrenó desde cero con 2014 a 2016. L0 se ajustó únicamente con δ50; se evaluó luego en los cinco deltas sin volver a entrenar. [B4]')
table(['Detector','Regla de selección','Configuración definitiva'],[
 ['L0','Mínima BCE de validación; máximo 60 épocas y paciencia 8','2 épocas · dropout 0.2'],
 ['V4','Mínima MSE de validación por estación; máximo 60 y paciencia 8','3 a 8 épocas · dropout 0.2'],
 ['PARAAD SA','Mínima MSE observada entre tres dropouts y 500 épocas','481 épocas · dropout 0.1']
],widths=[1.15,3.6,2.11])
h('Por qué dos épocas no equivalen a poco entrenamiento',2)
p('L0 ejecutó diez épocas durante selección. Su BCE de validación fue 0.023926 en la primera y 0.023046 en la segunda; no mejoró durante las ocho siguientes. La pérdida de ajuste sí continuó bajando, por lo que la regla prevista conservó la época 2. El definitivo se entrenó exactamente esas dos épocas.')
p('Una época es una pasada por el conjunto propio de cada modelo. L0 usó 546,500 ventanas iniciales y batch 64: aproximadamente 8,540 actualizaciones por época. PARAAD usó cientos de bloques diarios y batch 256: tres actualizaciones por época durante selección. También son objetivos distintos, clasificación y reconstrucción. No se igualan las épocas entre métodos ni se comparan sus pérdidas numéricas como si midieran lo mismo.')
table(['Dropout PARAAD','Época de mínimo','MSE limpia en 2016'],[
 ['0.1','481','0.174476'],['0.3','457','0.186258'],['0.5','481','0.175505']
],widths=[2.1,1.7,3.06])
h('Reproducibilidad',2)
p('Se realizaron 32 pares, 64 entrenamientos científicos, además de los smoke separados. En cada par coincidieron exactamente las historias, todos los pesos, las variables del optimizador y una predicción de control fija de hasta 256 ejemplos. Se comprobó la recarga del modelo guardado y sus hashes.')
p('Entorno CPU de un hilo, semilla 42 antes de construir y de entrenar, TensorFlow 2.21.0, Keras 3.13.2, NumPy 2.1.3 y Python 3.13.18. Esta igualdad está verificada dentro del mismo entorno; no supone identidad entre cualquier hardware o versión.')

# 7
page('Calibración y estabilidad de las falsas alarmas')
p('Los modelos congelados se cargaron de disco y se puntuaron dos veces sobre 2017 limpio, con 4,536 estación días comunes. No se reajustaron pesos ni épocas. Se congeló un umbral global por detector para p99 y p99.9, con alerta estricta puntaje > umbral. [B5]')
table(['Detector','Puntaje diario','p99','p99.9'],[
 ['L0','Probabilidad media','0.079395','0.356570'],
 ['V4','|Media firmada del residuo| / p95 propio','0.826410','1.115136'],
 ['PARAAD','MAE observado estandarizado','0.651824','0.818481']
],widths=[.9,3.36,1.3,1.3])
p('El p95 V4 se obtuvo una sola vez del residuo horario absoluto de 2017 por estación, como escala de puntaje, y se mantuvo fijo en ambos años de prueba. No modifica sus entradas ni su entrenamiento. Los tres umbrales numéricos tienen distintas escalas y no se comparan por su magnitud.')
p('Cada detector produjo 45 alertas de 4,536 casos con p99 (0.9921%) y cuatro con p99.9 (0.0882%). La resolución es una alerta entre 4,536. Estos porcentajes son empíricos de calibración; no garantizan la misma tasa en años posteriores.')
data=[]
for y in [2018,2019]:
    for d in ['l0','v4','paraad']:
        x={r['punto']:r for r in select(FA,ano=y,detector=d)}
        data.append([y,d.upper(),pct(x['p99']['tasa']),pct(x['p99.9']['tasa'])])
table(['Año limpio','Detector','Tasa p99','Tasa p99.9'],data,widths=[1.2,1.2,2.23,2.23])
p('En 2019 aumentaron las tasas limpias de los tres métodos. PARAAD llegó a 2.4743% con p99 y a 0.8102% con p99.9. Esta variación temporal debe acompañar cualquier afirmación sobre recall o F1; no se corrigió recalibrando sobre el examen.')
p('La tasa de esta tabla considera negativos experimentales todos los casos del año sin inyección. Es distinta del control “sin ataque” que conserva las etiquetas de los planes, presentado más adelante.')

# 8
page('Resultados principales de Barcelona con p99')
p('Se evaluaron 100 matrices alteradas y dos limpias por detector, más un control de 2017. Los 306 pares de inferencia de examen y tres pares de control fueron idénticos. El control de 2017 reprodujo exactamente los puntajes congelados: diferencia máxima cero y ningún cambio de alerta en los tres métodos. [B6]')
table(['Año','δ','L0 supervisado','V4','PARAAD'],f1rows('p99'),widths=[.7,.7,1.82,1.82,1.82])
p('F1 se obtiene sumando las confusiones de los diez calendarios de cada año e intensidad, no promediando sus F1. El mismo caso ambiental aparece en distintas realizaciones; la incertidumbre mantiene juntas esas repeticiones. Los deltas se reportan completos, sin seleccionar sólo los favorables.')
h('Qué ocurre con ataques del 50 por ciento',2)
data=[]
for y in [2018,2019]:
    for d in ['l0','v4','paraad']:
        r=met(y,50,d);data.append([y,d.upper(),f(r['precision']),f(r['recall']),f(r['f1'])])
table(['Año','Detector','Precisión','Recall','F1'],data,widths=[.7,1.0,1.72,1.72,1.72])
p('En 2018, L0 detectó aproximadamente 78 de cada 100 estación días atacados, frente a 39 de PARAAD y 26 de V4. Sin embargo, sólo alrededor de 37 de cada 100 alertas de L0 correspondieron a positivos del plan. En 2019 su recall se mantuvo cerca de 77%, pero la precisión bajó a 30%. La baja prevalencia y las falsas alarmas importan para interpretar la utilidad práctica.')
p('F1 no es porcentaje de aciertos. Combina precisión y recall: F1 = 2TP / (2TP + FP + FN). Un F1 de 0.50 no significa que el detector acierte en la mitad de todas las observaciones.')

# 9 Landscape figure
sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width=Inches(11.6929);sec.page_height=Inches(8.2677)
sec.top_margin=Inches(.55);sec.bottom_margin=Inches(.55);sec.left_margin=Inches(.55);sec.right_margin=Inches(.55)
h('Comparación visual del F1 en los dos años')
p('Barras con ataques proporcionales e intervalos del 95%. El mismo umbral p99 de 2017 se aplica a los dos años. Las líneas discontinuas muestran la referencia sin inyección con las mismas etiquetas programadas.')
q=doc.add_paragraph();q.paragraph_format.keep_with_next=True
r=q.add_run();r.add_picture(str(R/'f1_barcelona_p99.png'),width=Inches(10.5))
inline=r._r.xpath('.//wp:docPr')[0];inline.set('descr','F1 p99 por intensidad y detector para Barcelona 2018 y 2019 con intervalos de confianza y control sin inyección')
caption('Figura 1. X: cinco intensidades porcentuales discretas. Y: F1 por estación día. Barras: confusiones sumadas de diez calendarios. Intervalos: 2,000 remuestreos de bloques estación × mes, con realizaciones y detectores pareados. Modelos cargados de disco; L0 es supervisado. La prevalencia fue 0.8090% en 2018 y 0.9131% en 2019.')
p('La figura permite leer el patrón por magnitud y año. Para decidir si una diferencia está respaldada por incertidumbre deben consultarse los intervalos pareados, no sólo comparar visualmente barras o intervalos individuales.')

# 10
sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.PORTRAIT
sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
h('Diferencias pareadas e interpretación de la evidencia')
p('Los IC95% se obtuvieron con 2,000 bootstrap de 156 bloques estación × mes por año. Cada bloque conserva juntas sus diez realizaciones y los tres detectores. Se usan los mismos pesos de remuestreo en las diferencias; no se tratan los diez calendarios como muestras ambientales independientes. Son intervalos por contraste, sin corrección por multiplicidad.')
data=[]
for y in [2018,2019]:
    for delta in [100,50,20,10,5]:
        vals=[]
        for comp,sign in [('v4-paraad',-1),('l0-paraad',1),('l0-v4',1)]:
            x=select(D,ano=y,delta=delta,punto='p99',estado='con_ataque',comparacion=comp)[0]
            lo=float(x['inferior']) if sign==1 else -float(x['superior']);hi=float(x['superior']) if sign==1 else -float(x['inferior'])
            vals.append(f"{sign*float(x['diferencia']):+.4f}\n[{lo:+.4f}; {hi:+.4f}]")
        data.append([y,f'{delta}%',*vals])
table(['Año','δ','PARAAD menos V4','L0 menos PARAAD','L0 menos V4'],data,widths=[.6,.6,1.89,1.89,1.88],size=9.0)
h('Comparación principal sin etiquetas de ataque',2)
p('PARAAD tiene F1 puntual mayor que V4 en las cinco intensidades de ambos años. En 2018 todos los intervalos PARAAD−V4 incluyen cero: no se sostiene una ventaja concluyente en la prueba principal. En 2019 excluyen cero a favor de PARAAD para δ100, δ50, δ20 y δ10; δ5 sigue sin separación clara.')
h('Referencia supervisada L0',2)
p('Con p99, L0 supera a ambos en δ50 y δ20 con intervalos positivos en los dos años. En δ100 de 2018, PARAAD tiene mayor F1 puntual, pero L0−PARAAD incluye cero. En 2019 con δ100, L0 queda por encima de ambos con intervalos positivos. Esto no demuestra superioridad aislada de arquitectura, porque L0 recibió etiquetas de ataque al entrenar.')

# 11
page('Segundo umbral y controles sin inyección')
p('p99.9 fue un segundo punto registrado antes del examen. Sus resultados se conservan aunque sean favorables a un método; p99 sigue siendo el análisis principal. No se elige el umbral después de observar cuál produce el mayor F1.')
table(['Año','δ','L0 supervisado','V4','PARAAD'],f1rows('p99.9'),widths=[.7,.7,1.82,1.82,1.82])
p('L0 con δ50 alcanza F1 de 0.6019 y 0.6033 en 2018 y 2019 con p99.9. La precisión sube a 0.6964 y 0.6492, pero el recall baja a 0.5299 y 0.5635. El aumento de F1 refleja ese intercambio entre alertas y detección; no se sustituye la regla principal registrada.')
h('F1 del control sin ataque',2)
data=[]
for y in [2018,2019]:
    for point in ['p99','p99.9']:
        data.append([y,point,*[f(met(y,50,d,point,'sin_ataque')['f1']) for d in ['l0','v4','paraad']]])
table(['Año','Umbral','L0','V4','PARAAD'],data,widths=[.7,.9,1.75,1.75,1.76])
p('En este control los modelos puntúan los datos limpios y se mantienen las etiquetas programadas de cada calendario. Se miden coincidencias con el plan sin alterar la señal. Como las fechas se comparten entre deltas, el F1 de referencia es idéntico en las cinco intensidades de cada año y umbral.')
p('Con δ5 y δ10 los F1 absolutos son bajos, aproximadamente entre 0.007 y 0.027 con p99, y varios se aproximan al control. No se ha calculado un contraste adicional ataque−sin ataque; los intervalos pareados publicados comparan detectores dentro de cada condición. No se atribuye significación a pequeñas diferencias frente al control sin ese análisis.')

# 12
page('Antecedente completado de PARAAD en RAMA')
p('El primer brazo del estudio ya había adaptado PARAAD SA a RAMA en la revisión 20. Se recuperan aquí sus resultados aceptados, sin volver a leer las matrices de examen ni usar los años 2024 a 2026 para modificar decisiones. El modelo PARAAD se entrenó con observaciones sin inyección de 2020 a 2022 y se calibró con 2023 limpio. [R1]')
p('Se evaluaron el escenario original, una estación a la vez y un solo nodo comprometido, con los cinco deltas y dos umbrales. Hubo 765 matrices alteradas y tres limpias, con dos inferencias idénticas por matriz. La configuración congelada fue dropout 0.3 y 499 épocas; estos parámetros no se reutilizaron como elección óptima para Barcelona.')
table(['Año RAMA','L0 único','V4','PARAAD adaptado'],[
 ['2024','0.388','0.101','0.128'],['2025','0.445','0.108','0.208'],['2026 parcial','0.502','0.193','0.140']
],widths=[1.55,1.5,1.5,2.31])
caption('Tabla de contexto. F1 del escenario original RAMA, δ50 y p99, tomada del reporte cerrado de revisión 20. El año 2026 abarca enero a julio. No son datos de Barcelona ni un agregado entre ciudades.')
table(['Año RAMA','PARAAD menos V4','IC95%'],[
 ['2024','+0.027','[−0.005; +0.066]'],['2025','+0.100','[+0.054; +0.157]'],['2026 parcial','−0.053','[−0.119; +0.034]']
],widths=[1.55,2.1,3.21])
p('En RAMA no hubo una mejora estable de PARAAD frente a V4 entre años. En 2025, δ50 y p99, la diferencia favoreció a PARAAD con intervalo que excluyó cero; no ocurrió así en 2024 ni 2026 parcial. L0 tuvo mayor F1 que PARAAD en las 36 combinaciones de δ100/δ50, tres escenarios, tres años y dos umbrales reportadas.')
h('Qué aporta contar con las dos direcciones',2)
p('Los resultados permiten discutir tanto la adaptación del método publicado a RAMA como la aplicación de la metodología local en Barcelona. El contraste principal sin etiquetas se mantiene V4−PARAAD; L0 se identifica como referencia supervisada en ambos brazos.')
p('No se comparan directamente F1 entre ciudades para atribuir una causa urbana. Cambian unidades, periodos, número de estaciones, prevalencia, cobertura y escenarios. La ventaja de L0 en algunos ataques también incorpora su acceso a ejemplos etiquetados. Estos resultados no demuestran que el tamaño de una ciudad explique el rendimiento ni invalidan el trabajo original de PARAAD.')

# 13
page('Qué se puede afirmar y qué queda pendiente')
h('Resultados que sí están respaldados',2)
bullet('La comparación se completó en las dos direcciones, con modelos nuevos o congelados según correspondía, controles de reproducibilidad y trazabilidad de artefactos.')
bullet('En Barcelona, con p99, L0 mostró ventaja frente a ambos métodos para δ20 y δ50 en los dos años, respaldada por intervalos pareados positivos.')
bullet('Entre los métodos sin etiquetas de ataque, PARAAD fue mejor en estimación puntual que V4 en ambos años de Barcelona, con diferencias separadas de cero en varias condiciones de 2019.')
bullet('Las perturbaciones pequeñas siguen siendo difíciles y los umbrales limpios fijos no mantienen necesariamente la misma tasa de alerta en años posteriores.')
h('Límites de la conclusión',2)
p('La cohorte Barcelona es propia y contiene 13 estaciones. La división temporal, las máscaras, los bloques UTC y el umbral fijo limpio difieren del experimento exacto del artículo. Por ello, se reporta una comparación controlada con datos de la misma fuente y no la reproducción literal de sus cifras publicadas.')
p('La supervisión y las representaciones son distintas. No es válido convertir una ventaja práctica de L0 en una demostración causal de superioridad de arquitectura. Tampoco se infiere que un detector sea universalmente mejor para cualquier año, magnitud o umbral.')
p('El bootstrap por estación y mes preserva dependencia temporal dentro del bloque y junta las realizaciones, pero no resuelve toda la dependencia espacial o entre meses. Se reportan todos los contrastes previstos; no se presenta una prueba global corregida por multiplicidad.')
h('Redacción posible para la tesis',2)
p('La evaluación controlada en Barcelona aporta evidencia de que la ventaja del detector local supervisado ante alteraciones proporcionales moderadas no es exclusiva del escenario RAMA. Esta ventaja se observó para δ20 y δ50 con el umbral principal y una separación temporal independiente. Sin embargo, la supervisión impide atribuirla únicamente a la arquitectura y persisten dificultades ante perturbaciones pequeñas y variación temporal de las falsas alarmas.')
h('Trabajo que permanece separado',2)
p('El sesgo aditivo de PARAAD quedó registrado y sus calendarios están generados: cuatro de 13 estaciones, 30% de días elegibles y 125 planes por año de prueba. No se aplicó ni evaluó ese complemento en las tablas de este informe. Los planes aditivos iniciales de entrenamiento se conservan como historial y no se usaron para ajustar el L0 proporcional principal.')
p('El siguiente trabajo es revisar esta síntesis con el tutor, decidir el lugar del complemento aditivo en el alcance final y redactar la discusión integrada. La fase proporcional ya terminada no necesita otra calibración ni elegir nuevas épocas con el examen. No se hizo commit ni se modificaron resultados RAMA cerrados.')

# 14
page('Inventario y referencias para recuperar el trabajo')
p('Las rutas siguientes son relativas a la raíz de lorawan11-rama-detection. Los manifiestos guardan hashes completos; este documento resume los resultados y permite localizar su evidencia.')
table(['Evidencia','Ubicación en el repositorio'],[
 ['Datos y cohorte','data/reference/barcelona_paraad/'],
 ['Fuentes originales','data/raw/barcelona_paraad/'],
 ['Rejilla y ventanas','results/barcelona_paraad/tiempo_utc_v1/\nresults/barcelona_paraad/ventanas_v1/'],
 ['Planes y series alteradas','results/barcelona_paraad/planes_proporcionales_v1/\nresults/barcelona_paraad/ataques_proporcionales_v1/'],
 ['Modelos y selección','results/barcelona_paraad/colab_20261029/'],
 ['Calibración congelada','results/barcelona_paraad/calibracion_2017_colab/'],
 ['Examen y métricas','results/barcelona_paraad/examen_colab_2018_2019/\nSubcarpeta resumen con CSV figura y ficha'],
 ['Implementación y controles','experiments/barcelona/\ntests/test_barcelona_*.py'],
 ['Antecedente RAMA','results/anomalias_red_v1/revision20/']
],widths=[1.6,5.26],size=8.8)
h('Reportes internos utilizados',2)
for ident,path in [
 ('B1','docs/auditoria_barcelona_paraad_2014_2019.md'),('B2','docs/ventanas_barcelona_paraad_v1.md'),
 ('B3','docs/escenario_proporcional_barcelona_v1.md'),('B4','docs/resultados_entrenamiento_barcelona_v1.md'),
 ('B5','docs/resultados_calibracion_barcelona_2017.md'),('B6','docs/resultados_examen_barcelona_2018_2019.md'),
 ('R1','docs/resultados_examen_paraad_revision20.md')]:
    q=p(f'[{ident}] {path}');q.paragraph_format.space_after=Pt(3)
    for r in q.runs:r.font.size=Pt(9)
h('Fuentes públicas',2)
for text in [
 '[L1] Allka y colaboradores (2024). Pattern-Based Attention Recurrent Autoencoder for Anomaly Detection in Air Quality Sensor Networks. IEEE TNSE 11(6), 6372–6381. DOI 10.1109/TNSE.2024.3454459.',
 '[C1] Código público de PARAAD: github.com/XhensildaAllka/PARAAD. Implementación SA contrastada con el commit 6dbb5a9144a043d7825d63f49463a442ae60165b.',
 '[D1] Generalitat de Catalunya. Dataset tasf-thgu de calidad del aire. analisi.transparenciacatalunya.cat/resource/tasf-thgu.json. Descargas originales y consultas conservadas en el repositorio.',
 '[D2] Àrea Metropolitana de Barcelona. Municipios metropolitanos. amb.cat/en/web/area-metropolitana/municipis-metropolitans. Delimitación administrativa adoptada en la cohorte.'
]:
    q=p(text)
    for r in q.runs:r.font.size=Pt(9)

for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
for paragraph in doc.paragraphs:
    for border in paragraph._p.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.save(DEST)
fuentes=[R/'metricas.csv',R/'intervalos.csv',R/'diferencias_pareadas.csv',R/'tasa_limpia.csv',R/'verificacion_final.json',ROOT/'docs/resultados_examen_paraad_revision20.md',ROOT/'data/reference/barcelona_paraad/cohorte13.csv']
(OUT/'fuentes_informe_20261029.json').write_text(json.dumps(dict(documento=DEST.name,fecha='2026-10-29',fuentes={str(x.relative_to(ROOT)):hashlib.sha256(x.read_bytes()).hexdigest() for x in fuentes}),ensure_ascii=False,indent=2)+'\n')
print(DEST)
